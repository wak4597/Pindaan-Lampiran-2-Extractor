import io
import re
import streamlit as st
from docx import Document
from docx.shared import Pt
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.enum.text import WD_ALIGN_PARAGRAPH

# --- Formatting Constants ---
BOLD_TERMS = [
    r"diperakui dengan terma atau syarat",
    r"diperakui"
]

CUSTOM_ENGLISH_WORDS = {
    "accident", "activities", "adl", "administration", "administrator",
    "affiliated", "amp", "and", "angiografi", "area", "assessment",
    "association", "asthma", "avenue", "barrier", "based", "basement",
    "bay", "bays", "beam", "bed", "bedded", "bereavement", "block",
    "bobath", "brachytherapy", "breastfeeding", "bulk", "bunker",
    "business", "c-arm", "cafeteria", "care", "cashier", "catheterization",
    "centre", "chair", "changing", "chemotherapy", "children", "clean",
    "clinic", "clinical", "clinics", "closet", "cold", "collection",
    "community", "cone", "consumable", "corridor", "couch", "counselling",
    "counter", "crrt", "cssd", "cssu", "cytotoxic", "day", "daycare",
    "dayward", "dental", "department", "desk", "dialysis", "dining",
    "dirty", "disposal", "district", "doctor's", "double", "double-bedded",
    "drug", "east", "echo", "ecg", "eeg", "endoscopy", "entrance",
    "equipment", "examination", "exam", "executive", "existing",
    "extended", "eye", "facilities", "facillities", "female", "fertility",
    "filing", "first", "floor", "focused", "four",
    "four-bedded", "function", "gamma", "general", "ground", "gym",
    "gymnasium", "haemodialysis", "hand", "hdu", "health", "high",
    "holding", "home", "hospice", "hospital", "hot", "icu", "imaging",
    "inhouse", "intensity", "intermediate", "international", "intra",
    "isolation", "it", "janitor", "janitor's", "joint", "knife",
    "laboratory", "laser", "ldr", "level", "lift", "linear", "lobby",
    "lounge", "lower", "lung", "major", "male", "management", "manager",
    "maternity", "medical", "medication", "medicine", "meeting",
    "mezzanine", "minor", "mobile", "multidisiplin", "negative",
    "nephrologist", "nerve", "new", "non-hepatitis", "nurse", "nursery",
    "nursing", "observation", "occupational", "off", "office", "oncology",
    "open", "operation", "optometrist", "optometry", "oral", "orthopaedic",
    "outpatient", "outsourced", "palliative", "park", "patient", "pediatric",
    "pharmacy", "physical", "physio", "physiotherapy", "physician's",
    "plastic", "play", "positive", "prep", "preparation", "pressure",
    "private", "procedure", "processing", "public", "quarantine",
    "radiology", "radiologist", "reception", "reconstitution", "record",
    "recovery", "refuge", "registration", "rehab", "rehabilitation",
    "relokasi", "reporting", "resusitasi", "retail", "robotic", "room",
    "rooms", "sample", "satellite", "scanner", "scn", "scope", "screening",
    "second", "serum", "server", "services", "settlement", "single",
    "single-bedded", "sleep", "sluice", "society", "soiled", "special",
    "specialist", "specialists", "speech", "staff", "station",
    "sterilization", "storage", "store", "stretcher", "study", "suite",
    "support", "surgery", "surgical", "table", "tea", "technical", "test",
    "theatre", "therapy", "third", "three", "three-bedded", "toilet",
    "tower", "traction", "training", "treadmill", "treatment", "triage",
    "trolley", "ultrasound", "unit", "utility", "vacant", "vascular",
    "vertigo", "vestibular", "vision", "visiting", "volunteer", "waiting",
    "ward", "wheelchair", "wing", "women's", "x-ray"
}


# --- Extractor ---
def extract_field_data(doc):
    """Extracts fields using positional table coordinates and targeted regex."""
    data = {
        "nama_premis": "",
        "alamat_premis": "",
        "no_fail": "",
        "pegawai_proses": "",
        "isu_ulasan": ""
    }

    # --- Step 1: Build safe sequential text for Regex fields (No Fail & Pegawai) ---
    sequential_lines = []
    for element in doc.element.body.iterchildren():
        if isinstance(element, CT_P):
            p = Paragraph(element, doc)
            if p.text.strip():
                sequential_lines.append(p.text.strip())
        elif isinstance(element, CT_Tbl):
            table = Table(element, doc)
            for row in table.rows:
                row_cells = []
                seen_cells = set()
                for cell in row.cells:
                    if cell._tc not in seen_cells:
                        seen_cells.add(cell._tc)
                        if cell.text.strip():
                            row_cells.append(cell.text.strip())
                if row_cells:
                    sequential_lines.append(" | ".join(row_cells))
                    
    full_document_text = "\n".join(sequential_lines)

    # 2. No. Fail (Matches KKM87/A2/... with or without Jilid)
    fail_match = re.search(r'(KKM[\s/]*87/A2/[A-Za-z0-9/.-]+(?:[\s]*JILID[\s]*\d+)?)', full_document_text, re.IGNORECASE)
    if fail_match:
        data["no_fail"] = fail_match.group(1).strip()

    # 3. Pegawai Proses (Captures name directly, retaining Dr./KJ prefixes)
    pegawai_match = re.search(r'Pegawai\s*Proses[\s:\|]*([^\n\r\|]+)', full_document_text, re.IGNORECASE)
    if pegawai_match:
        data["pegawai_proses"] = pegawai_match.group(1).strip()

    # --- Step 2: Positional Extraction for Table Data ---
    for table in doc.tables:
        if len(table.rows) >= 2:
            
            # Safely get actual columns of the 2nd row (index 1), skipping merged cell duplicates
            row_2_cells = []
            seen_cells = set()
            for cell in table.rows[1].cells:
                if cell._tc not in seen_cells:
                    seen_cells.add(cell._tc)
                    row_2_cells.append(cell)
            
            # Ensure it's the right table (at least 4 columns)
            if len(row_2_cells) >= 4:
                
                # --- 1. Nama Premis & Alamat Premis (2nd Row, 2nd Column) ---
                cell_2_2_text = row_2_cells[1].text.strip()
                premis_match = re.search(r'([\s\S]*?)(?=Pemohon:)', cell_2_2_text, re.IGNORECASE)
                raw_premis_alamat = premis_match.group(1).strip() if premis_match else cell_2_2_text
                
                # Split Nama (First Line) and Alamat (Remaining Lines)
                lines = [line.strip() for line in raw_premis_alamat.split('\n') if line.strip()]
                if lines:
                    data["nama_premis"] = lines[0]
                    data["alamat_premis"] = " ".join(lines[1:])
                
                # --- 4. Isu / Ulasan ---
                
                # 4a. Keputusan (2nd Row, 4th Column -> Extract Bold Text Only)
                cell_2_4 = row_2_cells[3]
                keputusan_parts = []
                for p in cell_2_4.paragraphs:
                    for run in p.runs:
                        if run.bold and run.text.strip():
                            keputusan_parts.append(run.text.strip())
                keputusan = " ".join(keputusan_parts).strip()
                
                if not keputusan:  # Fallback if no bold formatting was natively used
                    keputusan = cell_2_4.text.split('\n')[0].strip()
                
                # 4b. Butiran Permohonan (2nd Row, 3rd Column)
                # Parse paragraph by paragraph to preserve list structure
                cell_2_3 = row_2_cells[2]
                cell_text_parts = []
                for p in cell_2_3.paragraphs:
                    text = p.text.strip()
                    if text:
                        style_name = p.style.name if p.style else ""
                        if 'Number' in style_name:
                            cell_text_parts.append(f"1. {text}")
                        elif 'List' in style_name or 'Bullet' in style_name:
                            cell_text_parts.append(f"• {text}")
                        else:
                            cell_text_parts.append(text)
                            
                cell_2_3_text = "\n".join(cell_text_parts)
                # Slice from "Permohonan:" up to "Ulasan/Pemerhatian:"
                butiran_match = re.search(r'Permohonan:\s*([\s\S]*?)(?=Ulasan\s*/\s*Pemerhatian:|$)', cell_2_3_text, re.IGNORECASE)
                butiran = butiran_match.group(1).strip() if butiran_match else cell_2_3_text

                # 4c. Tarikh Pindaan (End Row, 4th Column)
                end_row_cells = []
                seen_end = set()
                for cell in table.rows[-1].cells:
                    if cell._tc not in seen_end:
                        seen_end.add(cell._tc)
                        end_row_cells.append(cell)
                        
                tarikh = ""
                if len(end_row_cells) >= 4:
                    end_cell_text = end_row_cells[3].text.strip()
                    tarikh_match = re.search(r'(Tarikh pindaan bermula[\s\S]*)', end_cell_text, re.IGNORECASE)
                    tarikh = tarikh_match.group(1).strip() if tarikh_match else end_cell_text

                # Combine Keputusan, Butiran, and Tarikh into one block
                isu_parts = [keputusan, butiran, tarikh]
                data["isu_ulasan"] = "\n".join([p for p in isu_parts if p])
                
                break # Stop searching tables once the target is processed

    return data


# --- Builders & Formatters ---
def format_styled_text_run(paragraph, text):
    """Formats right-column text with Arial 9pt, selective bolding, and italics for English words."""
    if not text:
        return

    bold_pattern = f"({'|'.join(BOLD_TERMS)})"
    tokens = re.split(bold_pattern, text, flags=re.IGNORECASE)

    for token in tokens:
        if not token:
            continue
            
        is_bold = any(re.fullmatch(term, token, re.IGNORECASE) for term in BOLD_TERMS)
        sub_tokens = re.split(r'(\s+|[.,;:]+)', token)
        
        for sub in sub_tokens:
            if not sub:
                continue
            
            run = paragraph.add_run(sub)
            run.font.name = 'Arial'
            run.font.size = Pt(9)
            
            if is_bold:
                run.bold = True
            
            clean_word = re.sub(r'[^\w]', '', sub).lower()
            if clean_word in CUSTOM_ENGLISH_WORDS and not is_bold:
                run.italic = True


def build_summary_doc(records):
    """Creates in-memory Word Document formatted with Arial 9pt, selective bold labels, and styling."""
    doc = Document()
    
    # Title Section
    title_p = doc.add_paragraph()
    title_run = title_p.add_run('Laporan Semakan Permohonan - Summarized Table')
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(12)
    title_run.bold = True

    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    
    # Header Row
    hdr_cells = table.rows[0].cells
    hdr_titles = ["Maklumat Premis & Pegawai", "Isu / Ulasan"]
    
    for i, title in enumerate(hdr_titles):
        p = hdr_cells[i].paragraphs[0]
        run = p.add_run(title)
        run.font.name = 'Arial'
        run.font.size = Pt(9)
        run.bold = True

    # Populate Data
    for data in records:
        row_cells = table.add_row().cells
        
        # --- Left Column ---
        p_left = row_cells[0].paragraphs[0]
        
        # 1. Nama Premis (Bold, no label)
        val_premis = p_left.add_run(f"{data.get('nama_premis', '')}\n")
        val_premis.font.name = 'Arial'
        val_premis.font.size = Pt(9)
        val_premis.bold = True

        # 2. Alamat Premis (Value only, no label)
        val_alamat = p_left.add_run(f"{data.get('alamat_premis', '')}\n")
        val_alamat.font.name = 'Arial'
        val_alamat.font.size = Pt(9)
        val_alamat.bold = False

        # 3. No. Fail (Label + Value)
        lbl_fail = p_left.add_run("No. Fail: ")
        lbl_fail.font.name = 'Arial'
        lbl_fail.font.size = Pt(9)
        lbl_fail.bold = True
        
        val_fail = p_left.add_run(f"{data.get('no_fail', '')}\n")
        val_fail.font.name = 'Arial'
        val_fail.font.size = Pt(9)
        val_fail.bold = False

        # 4. Pegawai Proses (Label + Value)
        lbl_pegawai = p_left.add_run("Pegawai Proses: ")
        lbl_pegawai.font.name = 'Arial'
        lbl_pegawai.font.size = Pt(9)
        lbl_pegawai.bold = True
        
        val_pegawai = p_left.add_run(f"{data.get('pegawai_proses', '')}")
        val_pegawai.font.name = 'Arial'
        val_pegawai.font.size = Pt(9)
        val_pegawai.bold = False

        # --- Right Column ---
        right_text = data.get('isu_ulasan', '')
        lines = right_text.split('\n')
        
        first_paragraph_used = False
        for line in lines:
            clean_line = line.strip()
            if not clean_line:
                continue
            
            # Create a new paragraph line to prevent them from running together
            if not first_paragraph_used:
                p_right = row_cells[1].paragraphs[0]
                first_paragraph_used = True
            else:
                p_right = row_cells[1].add_paragraph()
            
            p_right.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            
            # Just print the line exactly as it was extracted
            format_styled_text_run(p_right, clean_line)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# --- Streamlit UI ---
st.set_page_config(page_title="Document Extractor", page_icon="📄", layout="centered")

st.title("📄 Pindaan to Lampiran 2 Extractor")
st.write("Upload your Word documents (.docx) below to extract data into a 2-column Word table.")

uploaded_files = st.file_uploader(
    "Choose Word files (.docx only)",
    type=["docx"],
    accept_multiple_files=True
)

if uploaded_files:
    if st.button("Process Documents", type="primary"):
        with st.spinner("Extracting data from files..."):
            extracted_records = []
            
            for file in uploaded_files:
                doc = Document(file)
                record = extract_field_data(doc)
                extracted_records.append(record)

            output_buffer = build_summary_doc(extracted_records)

        st.success("Extraction Complete!")
        
        st.download_button(
            label="📥 Download Extracted Summary (.docx)",
            data=output_buffer,
            file_name="Extracted_Summary.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
